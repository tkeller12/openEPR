import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# === CONFIG ===
filename = 'S-Band PFU Simple_1.0_BOM' # no extension
input_csv = filename + ".csv"
output_excel = filename + ".xlsx"
header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
alt_fill = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")
thin_border = Border(bottom=Side(style='thin'))

# === READ CSV ===
df = pd.read_csv(input_csv)

# === WRITE INITIAL EXCEL ===
df.to_excel(output_excel, index=False)

# === LOAD FOR FORMATTING ===
wb = load_workbook(output_excel)
ws = wb.active

# === FORMAT HEADERS ===
for col_num, col in enumerate(ws.iter_cols(min_row=1, max_row=1), 1):
    cell = col[0]
    cell.font = Font(bold=True)
    cell.fill = header_fill
    cell.border = thin_border

# === FORMAT ALTERNATING ROW COLORS ===
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    if (row[0].row % 2) == 0:
        for cell in row:
            cell.fill = alt_fill

# === AUTO-FIT COLUMN WIDTHS ===
for column_cells in ws.columns:
    max_length = 0
    column = column_cells[0].column_letter
    for cell in column_cells:
        try:
            cell_length = len(str(cell.value))
            if cell_length > max_length:
                max_length = cell_length
        except:
            pass
    adjusted_width = (max_length + 2)
    ws.column_dimensions[column].width = adjusted_width

# === SAVE ===
wb.save(output_excel)
print(f"Done! BOM exported to {output_excel} with formatting for easy assembly.")
